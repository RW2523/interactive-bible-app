#!/usr/bin/env python3
"""Build the Interactive Bible App demo media from demo_content/texts.

Usage (from the project root):
    backend/.venv/bin/python demo_content/build_demo_media.py [--force] [--jobs N]

Outputs:
    demo_content/media/sermon_all_things_for_good.mp4  narrated video (H.264 640x360, AAC)
    demo_content/media/sermon_all_things_for_good.vtt  one cue per sentence
    demo_content/media/podcast_forgiveness.mp3         two-voice narrated podcast
    demo_content/media/podcast_forgiveness.vtt         one cue per sentence, <v HOST>/<v GUEST>
    demo_content/docs/study_faith_hebrews11.pdf        fpdf2, core Helvetica font
    demo_content/docs/devotional_anxiety.docx          python-docx, real Heading 1/2 styles
    demo_content/manifest.json                         resource list for seeding

Requires macOS `say`, ffmpeg and ffprobe on PATH, plus fpdf2 and python-docx
(Pillow is optional; it draws the video title card).

Narration is synthesized one sentence at a time. Each clip is resampled to
22050 Hz mono PCM, measured with ffprobe (sample count), separated by fixed
silences and concatenated with the ffmpeg concat demuxer, so caption times are
exact sums of measured sample counts.

Steps whose outputs already exist are skipped unless --force is given. The
manifest is regenerated on every run from the media actually on disk.
"""

from __future__ import annotations

import argparse
import json
import re
import shutil
import subprocess
import sys
import tempfile
import wave
from concurrent.futures import ThreadPoolExecutor
from dataclasses import dataclass
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
DEMO = ROOT / "demo_content"
TEXTS = DEMO / "texts"
MEDIA = DEMO / "media"
DOCS = DEMO / "docs"
MANIFEST = DEMO / "manifest.json"

SAMPLE_RATE = 22050
SENTENCE_GAP_MS = 350  # silence between sentences
TURN_GAP_MS = 600  # silence when the podcast speaker changes
TAIL_MS = 250  # silence after the last sentence
SPEECH_RATE_WPM = 170
MAX_CUE_END_DRIFT_MS = 500

VIDEO_W, VIDEO_H, VIDEO_FPS = 640, 360, 15
VIDEO_BG_HEX = "0f172a"
WAVE_HEX = "7dd3fc"

# Novelty voices that should never be used for narration.
NOVELTY_VOICES = {
    "Albert", "Bad News", "Bahh", "Bells", "Boing", "Bubbles", "Cellos", "Good News",
    "Jester", "Junior", "Organ", "Superstar", "Trinoids", "Whisper", "Wobble", "Zarvox",
}

SERMON = {
    "key": "sermon_all_things_for_good",
    "title": "All Things for Good",
    "speaker": "Pastor Elena Brooks",
    "text": TEXTS / "sermon_all_things_for_good.txt",
    "video": MEDIA / "sermon_all_things_for_good.mp4",
    "vtt": MEDIA / "sermon_all_things_for_good.vtt",
}
PODCAST = {
    "key": "podcast_forgiveness",
    "title": "Grace Notes: Learning to Forgive",
    "speaker": "Maya Chen with Pastor Owen Hartley",
    "text": TEXTS / "podcast_forgiveness.txt",
    "audio": MEDIA / "podcast_forgiveness.mp3",
    "vtt": MEDIA / "podcast_forgiveness.vtt",
}
STUDY = {
    "key": "study_faith_hebrews11",
    "title": "Faith That Endures: A Small Group Study on Hebrews 11",
    "author": "Marcus Bell",
    "md": TEXTS / "study_faith_hebrews11.md",
    "pdf": DOCS / "study_faith_hebrews11.pdf",
}
DEVOTIONAL = {
    "key": "devotional_anxiety",
    "title": "Peace for Anxious Hearts",
    "author": "Clara Whitfield",
    "md": TEXTS / "devotional_anxiety.md",
    "docx": DOCS / "devotional_anxiety.docx",
}
ARTICLE = {
    "key": "article_prayer",
    "title": "Learning to Pray Again",
    "author": "Hannah Okoro",
    "md": TEXTS / "article_prayer.md",
}


# --------------------------------------------------------------------------- helpers


def rel(path: Path) -> str:
    return path.relative_to(ROOT).as_posix()


def run(cmd: list, capture: bool = False) -> str:
    args = [str(c) for c in cmd]
    res = subprocess.run(args, capture_output=True, text=True)
    if res.returncode != 0:
        raise RuntimeError(
            f"command failed (exit {res.returncode}): {' '.join(args)}\n{res.stderr.strip()}"
        )
    return res.stdout if capture else ""


def ffprobe(path: Path, entries: str) -> dict:
    out = run(["ffprobe", "-v", "error", "-show_entries", entries, "-of", "json", path], capture=True)
    return json.loads(out)


def media_duration_ms(path: Path) -> int:
    return round(float(ffprobe(path, "format=duration")["format"]["duration"]) * 1000)


def wav_samples(path: Path) -> int:
    """Exact sample count of a mono 22050 Hz PCM WAV, measured with ffprobe."""
    info = ffprobe(path, "stream=sample_rate,channels,duration_ts,time_base")["streams"][0]
    if int(info["sample_rate"]) != SAMPLE_RATE or int(info["channels"]) != 1:
        raise RuntimeError(f"{path.name}: expected {SAMPLE_RATE} Hz mono, got {info}")
    if info.get("time_base") != f"1/{SAMPLE_RATE}":
        raise RuntimeError(f"{path.name}: unexpected time_base {info.get('time_base')}")
    samples = int(info["duration_ts"])
    with wave.open(str(path), "rb") as w:  # cross-check against the RIFF header
        if w.getnframes() != samples:
            raise RuntimeError(f"{path.name}: ffprobe={samples} samples, header={w.getnframes()}")
    return samples


def samples_to_ms(samples: int) -> int:
    return (samples * 1000 + SAMPLE_RATE // 2) // SAMPLE_RATE


def fmt_ts(ms: int) -> str:
    hours, rem = divmod(ms, 3_600_000)
    minutes, rem = divmod(rem, 60_000)
    seconds, millis = divmod(rem, 1000)
    return f"{hours:02d}:{minutes:02d}:{seconds:02d}.{millis:03d}"


def human_size(path: Path) -> str:
    size = path.stat().st_size
    return f"{size / 1_048_576:.2f} MB" if size >= 1_048_576 else f"{size / 1024:.1f} KB"


def check_tools() -> None:
    missing = [tool for tool in ("say", "ffmpeg", "ffprobe") if shutil.which(tool) is None]
    if missing:
        sys.exit(f"error: required tool(s) not found on PATH: {', '.join(missing)}")


# --------------------------------------------------------------------------- text handling

_ABBREVIATIONS = {"mr.", "mrs.", "ms.", "dr.", "st.", "rev.", "vs.", "etc.", "e.g.", "i.e.", "ch.", "v.", "vv."}
_SENTENCE_END = re.compile(r"[.!?]+[\"'”’)\]]*(?=\s+[\"'“‘(\[]?[A-Z0-9])")


def split_sentences(text: str) -> list[str]:
    text = " ".join(text.split())
    sentences, start = [], 0
    for match in _SENTENCE_END.finditer(text):
        chunk = text[start:match.end()].strip()
        if chunk.rsplit(" ", 1)[-1].lower() in _ABBREVIATIONS:
            continue
        sentences.append(chunk)
        start = match.end()
    tail = text[start:].strip()
    if tail:
        sentences.append(tail)
    return sentences


def tts_text(sentence: str) -> str:
    """Spoken form of a caption sentence (captions keep the original wording)."""
    sentence = sentence.replace("LORD", "Lord")
    return sentence.translate({0x201C: '"', 0x201D: '"', 0x2018: "'", 0x2019: "'", 0x2014: ", ", 0x2013: "-"})


@dataclass
class Segment:
    text: str
    speaker: str


@dataclass
class Cue:
    start: int  # samples
    end: int  # samples
    text: str
    speaker: str


def sermon_segments() -> list[Segment]:
    lines = SERMON["text"].read_text(encoding="utf-8").splitlines()
    return [Segment(s, "PASTOR") for line in lines if line.strip() for s in split_sentences(line)]


def podcast_segments() -> list[Segment]:
    segments = []
    for number, line in enumerate(PODCAST["text"].read_text(encoding="utf-8").splitlines(), 1):
        if not line.strip():
            continue
        match = re.match(r"^(HOST|GUEST):\s*(.+)$", line)
        if not match:
            raise ValueError(f"{rel(PODCAST['text'])}:{number}: line must start with HOST: or GUEST:")
        segments.extend(Segment(s, match.group(1)) for s in split_sentences(match.group(2)))
    return segments


# --------------------------------------------------------------------------- narration


def installed_voices() -> dict[str, str]:
    """Voice name -> locale, parsed from `say -v '?'`."""
    out = run(["say", "-v", "?"], capture=True)
    voices = {}
    for line in out.splitlines():
        match = re.match(r"^(?P<name>.+?)\s+(?P<locale>[a-z]{2,3}[_-][A-Za-z0-9]+)\s+#", line)
        if match:
            voices[match.group("name").strip()] = match.group("locale")
    return voices


def pick_voice(voices: dict[str, str], preferred: list[str], exclude: set[str] = frozenset(),
               any_english: bool = False) -> str | None:
    """First preferred voice that exists; optionally any natural English voice; else None (system default)."""
    for name in preferred:
        if name in voices and name not in exclude:
            return name
    if any_english:
        for name, locale in sorted(voices.items(), key=lambda kv: (not kv[1].startswith("en_US"), kv[0])):
            if locale.startswith("en") and name not in NOVELTY_VOICES and name not in exclude:
                return name
    return None


def synthesize(index: int, segment: Segment, voice: str | None, workdir: Path) -> tuple[Path, int]:
    txt = workdir / f"s{index:04d}.txt"
    aiff = workdir / f"s{index:04d}.aiff"
    wav = workdir / f"s{index:04d}.wav"
    txt.write_text(tts_text(segment.text), encoding="utf-8")
    cmd = ["say", "-r", SPEECH_RATE_WPM, "-o", aiff, "-f", txt]
    if voice:
        cmd[1:1] = ["-v", voice]
    run(cmd)
    run(["ffmpeg", "-v", "error", "-y", "-i", aiff, "-ac", "1", "-ar", SAMPLE_RATE, "-c:a", "pcm_s16le", wav])
    txt.unlink()
    aiff.unlink()
    return wav, wav_samples(wav)


def write_silence(path: Path, ms: int) -> int:
    samples = round(SAMPLE_RATE * ms / 1000)
    with wave.open(str(path), "wb") as w:
        w.setnchannels(1)
        w.setsampwidth(2)
        w.setframerate(SAMPLE_RATE)
        w.writeframes(b"\x00\x00" * samples)
    return samples


def narrate(segments: list[Segment], voices: dict[str, str | None], workdir: Path, jobs: int) -> tuple[Path, list[Cue]]:
    """Synthesize every segment, join with silences, return (wav, cues with exact sample offsets)."""
    with ThreadPoolExecutor(max_workers=max(1, jobs)) as pool:
        clips = list(pool.map(lambda item: synthesize(item[0], item[1], voices[item[1].speaker], workdir),
                              enumerate(segments)))

    silences: dict[int, tuple[Path, int]] = {}

    def silence(ms: int) -> tuple[Path, int]:
        if ms not in silences:
            path = workdir / f"silence_{ms}ms.wav"
            silences[ms] = (path, write_silence(path, ms))
        return silences[ms]

    playlist, cues, offset = [], [], 0
    for i, (segment, (clip, samples)) in enumerate(zip(segments, clips)):
        cues.append(Cue(offset, offset + samples, segment.text, segment.speaker))
        playlist.append(clip)
        offset += samples
        if i + 1 < len(segments):
            gap_ms = TURN_GAP_MS if segments[i + 1].speaker != segment.speaker else SENTENCE_GAP_MS
        else:
            gap_ms = TAIL_MS
        gap_path, gap_samples = silence(gap_ms)
        playlist.append(gap_path)
        offset += gap_samples

    concat_list = workdir / "concat.txt"
    concat_list.write_text("".join(f"file '{p.as_posix()}'\n" for p in playlist), encoding="utf-8")
    joined = workdir / "narration.wav"
    run(["ffmpeg", "-v", "error", "-y", "-f", "concat", "-safe", "0", "-i", concat_list, "-c", "copy", joined])
    total = wav_samples(joined)
    if total != offset:
        raise RuntimeError(f"concatenated narration has {total} samples, timeline expects {offset}")
    for clip, _ in clips:
        clip.unlink()
    return joined, cues


def write_vtt(path: Path, cues: list[Cue], voice_tags: bool) -> None:
    def escape(text: str) -> str:
        return text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")

    blocks = ["WEBVTT", ""]
    for cue in cues:
        body = escape(cue.text)
        if voice_tags:
            body = f"<v {cue.speaker}>{body}"
        blocks += [f"{fmt_ts(samples_to_ms(cue.start))} --> {fmt_ts(samples_to_ms(cue.end))}", body, ""]
    path.write_text("\n".join(blocks), encoding="utf-8")


# --------------------------------------------------------------------------- sermon video


def title_card(path: Path, title: str, subtitle: str) -> Path | None:
    """Dark solid background with a title; returns None when Pillow is unavailable."""
    try:
        from PIL import Image, ImageDraw, ImageFont
    except ImportError:
        return None

    def font(size: int, bold: bool = False):
        candidates = [
            "/System/Library/Fonts/Supplemental/Arial Bold.ttf" if bold else "/System/Library/Fonts/Supplemental/Arial.ttf",
            "/System/Library/Fonts/Helvetica.ttc",
            "/Library/Fonts/Arial.ttf",
        ]
        for candidate in candidates:
            try:
                return ImageFont.truetype(candidate, size)
            except OSError:
                continue
        return ImageFont.load_default()

    bg = tuple(int(VIDEO_BG_HEX[i:i + 2], 16) for i in (0, 2, 4))
    image = Image.new("RGB", (VIDEO_W, VIDEO_H), bg)
    draw = ImageDraw.Draw(image)
    for text, size, bold, y, color in (
        (title, 30, True, 40, (241, 245, 249)),
        (subtitle, 16, False, 82, (148, 163, 184)),
        ("Interactive Bible App demo content · synthetic narration", 12, False, VIDEO_H - 28, (100, 116, 139)),
    ):
        f = font(size, bold)
        draw.text(((VIDEO_W - draw.textlength(text, font=f)) / 2, y), text, font=f, fill=color)
    image.save(path)
    return path


def encode_video(audio: Path, card: Path | None, out: Path, title: str, artist: str) -> None:
    if card is not None:
        background = ["-loop", "1", "-framerate", VIDEO_FPS, "-i", card]
    else:
        background = ["-f", "lavfi", "-i", f"color=c=0x{VIDEO_BG_HEX}:s={VIDEO_W}x{VIDEO_H}:r={VIDEO_FPS}"]
    graph = (
        f"[0:a]showwaves=s={VIDEO_W}x120:mode=cline:rate={VIDEO_FPS}:colors=0x{WAVE_HEX}:draw=full[w];"
        f"[1:v][w]overlay=0:130:shortest=1,format=yuv420p[v]"
    )
    run([
        "ffmpeg", "-v", "error", "-y", "-i", audio, *background,
        "-filter_complex", graph, "-map", "[v]", "-map", "0:a",
        "-c:v", "libx264", "-preset", "slow", "-crf", "34", "-maxrate", "64k", "-bufsize", "128k",
        "-g", VIDEO_FPS * 10, "-r", VIDEO_FPS, "-pix_fmt", "yuv420p",
        "-c:a", "aac", "-b:a", "64k", "-ac", "1",
        "-metadata", f"title={title}", "-metadata", f"artist={artist}",
        "-metadata", "comment=Interactive Bible App demo content; synthetic TTS narration",
        "-movflags", "+faststart", "-shortest", out,
    ])


def build_sermon(force: bool, jobs: int, voices: dict[str, str]) -> None:
    video, vtt = SERMON["video"], SERMON["vtt"]
    if not force and video.exists() and vtt.exists():
        print(f"skip   {rel(video)} + .vtt (exist; use --force to rebuild)")
        return
    segments = sermon_segments()
    voice = pick_voice(voices, ["Samantha"])
    print(f"build  {rel(video)}: {len(segments)} sentences, voice={voice or 'system default'}")
    with tempfile.TemporaryDirectory(prefix="ibible_sermon_") as tmp:
        workdir = Path(tmp)
        narration, cues = narrate(segments, {"PASTOR": voice}, workdir, jobs)
        card = title_card(workdir / "card.png", SERMON["title"], f"Sermon · {SERMON['speaker']}")
        staged = workdir / video.name
        encode_video(narration, card, staged, SERMON["title"], SERMON["speaker"])
        MEDIA.mkdir(parents=True, exist_ok=True)
        shutil.move(str(staged), video)
        write_vtt(vtt, cues, voice_tags=False)


# --------------------------------------------------------------------------- podcast audio


def build_podcast(force: bool, jobs: int, voices: dict[str, str]) -> None:
    audio, vtt = PODCAST["audio"], PODCAST["vtt"]
    if not force and audio.exists() and vtt.exists():
        print(f"skip   {rel(audio)} + .vtt (exist; use --force to rebuild)")
        return
    segments = podcast_segments()
    host = pick_voice(voices, ["Samantha"])
    guest_preferred = ["Daniel", "Alex", "Tom", "Aaron", "Arthur", "Oliver", "Rishi",
                       "Reed (English (US))", "Eddy (English (US))", "Rocko (English (US))"]
    guest = pick_voice(voices, guest_preferred, exclude={host} if host else set(), any_english=True)
    print(f"build  {rel(audio)}: {len(segments)} sentences, HOST={host or 'system default'}, "
          f"GUEST={guest or 'system default'}")
    with tempfile.TemporaryDirectory(prefix="ibible_podcast_") as tmp:
        workdir = Path(tmp)
        narration, cues = narrate(segments, {"HOST": host, "GUEST": guest}, workdir, jobs)
        staged = workdir / audio.name
        run([
            "ffmpeg", "-v", "error", "-y", "-i", narration,
            "-c:a", "libmp3lame", "-b:a", "64k", "-ac", "1", "-id3v2_version", "3",
            "-metadata", f"title={PODCAST['title']}", "-metadata", f"artist={PODCAST['speaker']}",
            "-metadata", "comment=Interactive Bible App demo content; synthetic TTS narration",
            staged,
        ])
        MEDIA.mkdir(parents=True, exist_ok=True)
        shutil.move(str(staged), audio)
        write_vtt(vtt, cues, voice_tags=True)


# --------------------------------------------------------------------------- documents


def markdown_blocks(markdown: str) -> list[tuple[str, str]]:
    """Tiny block parser: h1/h2/h3, ol, ul, quote, p (consecutive lines join into one paragraph)."""
    blocks: list[tuple[str, str]] = []
    paragraph: list[str] = []

    def flush() -> None:
        if paragraph:
            blocks.append(("p", " ".join(paragraph)))
            paragraph.clear()

    for raw in markdown.splitlines():
        line = raw.strip()
        heading = re.match(r"^(#{1,3})\s+(.*)$", line)
        ordered = re.match(r"^(\d+)[.)]\s+(.*)$", line)
        bullet = re.match(r"^[-*+]\s+(.*)$", line)
        if not line:
            flush()
        elif heading:
            flush()
            blocks.append((f"h{len(heading.group(1))}", heading.group(2).strip()))
        elif ordered:
            flush()
            blocks.append(("ol", f"{ordered.group(1)}.\t{ordered.group(2)}"))
        elif bullet:
            flush()
            blocks.append(("ul", bullet.group(1)))
        elif line.startswith(">"):
            flush()
            blocks.append(("quote", line.lstrip("> ").strip()))
        else:
            paragraph.append(line)
    flush()
    return blocks


def strip_inline_markdown(text: str) -> str:
    text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
    text = re.sub(r"(?<!\w)[*_](.+?)[*_](?!\w)", r"\1", text)
    return text.replace("`", "")


def latin1(text: str) -> str:
    """Core PDF fonts are latin-1 only: map typographic punctuation to ASCII."""
    table = {0x2018: "'", 0x2019: "'", 0x201C: '"', 0x201D: '"', 0x2013: "-", 0x2014: " - ",
             0x2026: "...", 0x00A0: " ", 0x00B7: "-", 0x2022: "-"}
    return text.translate(table).encode("latin-1", "replace").decode("latin-1")


def build_pdf(force: bool) -> None:
    out = STUDY["pdf"]
    if not force and out.exists():
        print(f"skip   {rel(out)} (exists; use --force to rebuild)")
        return
    try:
        from fpdf import FPDF
        from fpdf.enums import XPos, YPos
    except ImportError:
        sys.exit("error: fpdf2 is not installed; run with backend/.venv/bin/python")

    pdf = FPDF(format="Letter")
    pdf.set_margins(left=22, top=20, right=22)
    pdf.set_auto_page_break(auto=True, margin=20)
    pdf.set_title(latin1(STUDY["title"]))
    pdf.set_author(latin1(STUDY["author"]))
    pdf.set_subject("Interactive Bible App demo content: small group Bible study")
    pdf.set_creator("demo_content/build_demo_media.py")
    pdf.add_page()
    body_h = 5.8
    flow = {"align": "L", "new_x": XPos.LMARGIN, "new_y": YPos.NEXT}
    for kind, text in markdown_blocks(STUDY["md"].read_text(encoding="utf-8")):
        text = latin1(strip_inline_markdown(text))
        if kind == "h1":
            pdf.set_font("Helvetica", "B", 19)
            pdf.multi_cell(0, 9, text, **flow)
            pdf.set_font("Helvetica", "", 10)
            pdf.set_text_color(100, 116, 139)
            pdf.multi_cell(0, 5, latin1(f"By {STUDY['author']}"), **flow)
            pdf.set_text_color(0, 0, 0)
            pdf.ln(4)
        elif kind in ("h2", "h3"):
            if pdf.get_y() > pdf.h - pdf.b_margin - 35:  # keep a heading with the text after it
                pdf.add_page()
            else:
                pdf.ln(3)
            pdf.set_font("Helvetica", "B", 14 if kind == "h2" else 12)
            pdf.multi_cell(0, 7, text, **flow)
            pdf.ln(1.5)
        elif kind in ("ol", "ul"):
            marker, _, item = text.partition("\t") if kind == "ol" else ("-", "", text)
            pdf.set_font("Helvetica", "", 11)
            pdf.set_x(pdf.l_margin + 4)
            pdf.cell(7, body_h, marker)
            pdf.multi_cell(0, body_h, item, **flow)
            pdf.ln(1.2)
        elif kind == "quote":
            pdf.set_font("Helvetica", "I", 11)
            pdf.set_x(pdf.l_margin + 8)
            pdf.multi_cell(0, body_h, text, **flow)
            pdf.ln(2.5)
        else:
            pdf.set_font("Helvetica", "", 11)
            pdf.multi_cell(0, body_h, text, **flow)
            pdf.ln(2.5)
    DOCS.mkdir(parents=True, exist_ok=True)
    pdf.output(str(out))
    print(f"build  {rel(out)}")


def add_inline_runs(paragraph, text: str) -> None:
    for part in re.split(r"(\*\*.+?\*\*|(?<!\w)\*.+?\*(?!\w))", text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            paragraph.add_run(part[1:-1]).italic = True
        else:
            paragraph.add_run(part)


def build_docx(force: bool) -> None:
    out = DEVOTIONAL["docx"]
    if not force and out.exists():
        print(f"skip   {rel(out)} (exists; use --force to rebuild)")
        return
    try:
        from docx import Document
    except ImportError:
        sys.exit("error: python-docx is not installed; run with backend/.venv/bin/python")

    document = Document()
    props = document.core_properties
    props.title = DEVOTIONAL["title"]
    props.author = DEVOTIONAL["author"]
    props.subject = "Interactive Bible App demo content: devotional"
    props.comments = "Generated by demo_content/build_demo_media.py"
    for kind, text in markdown_blocks(DEVOTIONAL["md"].read_text(encoding="utf-8")):
        if kind in ("h1", "h2", "h3"):
            document.add_heading(strip_inline_markdown(text), level=int(kind[1]))  # "Heading N" style
        elif kind == "ul":
            add_inline_runs(document.add_paragraph(style="List Bullet"), text)
        elif kind == "ol":
            add_inline_runs(document.add_paragraph(style="List Number"), text.partition("\t")[2])
        elif kind == "quote":
            add_inline_runs(document.add_paragraph(style="Quote"), text)
        else:
            add_inline_runs(document.add_paragraph(), text)
    DOCS.mkdir(parents=True, exist_ok=True)
    document.save(str(out))
    print(f"build  {rel(out)}")


# --------------------------------------------------------------------------- manifest


def manifest_entries() -> list[dict]:
    def entry(**fields) -> dict:
        base = {
            "key": None, "title": None, "type": None, "category": None, "speaker": None, "author": None,
            "language": "en", "visibility": "public", "rights_status": "owned", "allow_clip_export": True,
            "is_official": False, "source_path": None, "captions_path": None, "body_text_path": None,
            "duration_ms": None, "description": None, "topic_hints": [],
        }
        base.update(fields)
        return base

    def duration(path: Path) -> int | None:
        return media_duration_ms(path) if path.exists() else None

    return [
        entry(
            key=SERMON["key"], title=SERMON["title"], type="video", category="sermon",
            speaker=SERMON["speaker"], is_official=True,
            source_path=rel(SERMON["video"]), captions_path=rel(SERMON["vtt"]),
            duration_ms=duration(SERMON["video"]),
            description="A narrated Sunday sermon on finding hope and God's providence in seasons of suffering.",
            topic_hints=["suffering", "hope", "providence", "perseverance", "trust"],
        ),
        entry(
            key=PODCAST["key"], title=PODCAST["title"], type="audio", category="podcast",
            speaker=PODCAST["speaker"],
            source_path=rel(PODCAST["audio"]), captions_path=rel(PODCAST["vtt"]),
            duration_ms=duration(PODCAST["audio"]),
            description="A two-voice podcast conversation about why forgiveness is hard and how to begin.",
            topic_hints=["forgiveness", "grace", "mercy", "reconciliation", "prayer"],
        ),
        entry(
            key=DEVOTIONAL["key"], title=DEVOTIONAL["title"], type="document", category="devotional",
            author=DEVOTIONAL["author"], source_path=rel(DEVOTIONAL["docx"]),
            description="A short devotional for anxious days: bringing worries to God and living one day at a time.",
            topic_hints=["anxiety", "peace", "trust", "prayer", "worry"],
        ),
        entry(
            key=STUDY["key"], title=STUDY["title"], type="pdf", category="study",
            author=STUDY["author"], source_path=rel(STUDY["pdf"]),
            description="A three-session small group Bible study on what faith is and how it endures testing.",
            topic_hints=["faith", "trust", "testing", "obedience", "promises of God"],
        ),
        entry(
            key=ARTICLE["key"], title=ARTICLE["title"], type="article", category="article",
            author=ARTICLE["author"], body_text_path=rel(ARTICLE["md"]),
            description="An article for anyone whose prayer life feels awkward or dry.",
            topic_hints=["prayer", "intercession", "spiritual disciplines", "humility"],
        ),
    ]


def write_manifest() -> None:
    content = json.dumps(manifest_entries(), indent=2, ensure_ascii=False) + "\n"
    if MANIFEST.exists() and MANIFEST.read_text(encoding="utf-8") == content:
        print(f"keep   {rel(MANIFEST)} (unchanged)")
        return
    MANIFEST.write_text(content, encoding="utf-8")
    print(f"write  {rel(MANIFEST)}")


# --------------------------------------------------------------------------- verification

_TIMING = re.compile(r"^(\d{2}):(\d{2}):(\d{2})\.(\d{3}) --> (\d{2}):(\d{2}):(\d{2})\.(\d{3})\s*$")


def parse_vtt(path: Path) -> list[tuple[int, int, str]]:
    lines = path.read_text(encoding="utf-8").splitlines()
    if not lines or not lines[0].startswith("WEBVTT"):
        raise ValueError(f"{rel(path)}: missing WEBVTT header")
    cues, i = [], 1
    while i < len(lines):
        match = _TIMING.match(lines[i])
        if not match:
            i += 1
            continue
        g = [int(x) for x in match.groups()]
        start = ((g[0] * 60 + g[1]) * 60 + g[2]) * 1000 + g[3]
        end = ((g[4] * 60 + g[5]) * 60 + g[6]) * 1000 + g[7]
        i += 1
        text = []
        while i < len(lines) and lines[i].strip():
            text.append(lines[i])
            i += 1
        cues.append((start, end, "\n".join(text)))
    return cues


def verify_media(media: Path, vtt: Path) -> list[str]:
    problems = []
    if not media.exists() or not vtt.exists():
        return [f"missing {rel(media) if not media.exists() else rel(vtt)}"]
    cues = parse_vtt(vtt)
    duration = media_duration_ms(media)
    if not cues:
        problems.append(f"{rel(vtt)}: no cues")
    previous_end = 0
    for n, (start, end, text) in enumerate(cues, 1):
        if not text.strip():
            problems.append(f"{rel(vtt)} cue {n}: empty text")
        if end <= start:
            problems.append(f"{rel(vtt)} cue {n}: end <= start")
        if start < previous_end:
            problems.append(f"{rel(vtt)} cue {n}: overlaps previous cue")
        previous_end = end
    drift = abs(duration - cues[-1][1]) if cues else 0
    if drift > MAX_CUE_END_DRIFT_MS:
        problems.append(f"{rel(vtt)}: last cue ends {cues[-1][1]} ms but media lasts {duration} ms")
    streams = ffprobe(media, "stream=codec_type,codec_name,width,height,r_frame_rate,sample_rate,channels,bit_rate")
    desc = ", ".join(
        f"{s['codec_type']}={s['codec_name']}"
        + (f" {s['width']}x{s['height']}@{s['r_frame_rate']}" if s["codec_type"] == "video" else f" {s.get('sample_rate')}Hz/{s.get('channels')}ch")
        for s in streams["streams"]
    )
    print(f"  {rel(media)}: {duration / 1000:.3f} s, {human_size(media)}, {desc}")
    print(f"  {rel(vtt)}: {len(cues)} cues, last cue end {cues[-1][1] / 1000 if cues else 0:.3f} s "
          f"(drift {drift} ms), monotonic={'yes' if not problems else 'NO'}")
    return problems


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--force", action="store_true", help="rebuild outputs even if they already exist")
    parser.add_argument("--jobs", type=int, default=4, help="parallel `say` processes (default 4)")
    args = parser.parse_args()

    check_tools()
    voices = installed_voices()
    build_sermon(args.force, args.jobs, voices)
    build_podcast(args.force, args.jobs, voices)
    build_pdf(args.force)
    build_docx(args.force)
    write_manifest()

    print("\nsummary")
    problems = verify_media(SERMON["video"], SERMON["vtt"]) + verify_media(PODCAST["audio"], PODCAST["vtt"])
    for path in (STUDY["pdf"], DEVOTIONAL["docx"], MANIFEST):
        print(f"  {rel(path)}: {human_size(path)}")
    total = sum(p.stat().st_size for d in (MEDIA, DOCS) for p in d.iterdir() if p.is_file())
    print(f"  media + docs total: {total / 1_048_576:.2f} MB")
    for problem in problems:
        print(f"PROBLEM: {problem}", file=sys.stderr)
    return 1 if problems else 0


if __name__ == "__main__":
    sys.exit(main())
